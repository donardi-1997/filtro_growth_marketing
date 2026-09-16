import importlib
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from marketing_cv_ranker import (
    phrase_hits,
    estimate_years_experience,
    score_candidate,
    classify_score,
)


def test_modular_package_boundaries_and_compatibility_exports():
    module_names = [
        "growth_ranker.profile",
        "growth_ranker.text",
        "growth_ranker.documents",
        "growth_ranker.experience",
        "growth_ranker.scoring",
        "growth_ranker.exporting",
        "growth_ranker.dashboard",
        "growth_ranker.gui",
        "growth_ranker.cli",
    ]
    modules = {name: importlib.import_module(name) for name in module_names}

    import marketing_cv_ranker as facade

    assert facade.phrase_hits is modules["growth_ranker.text"].phrase_hits
    assert facade.estimate_years_experience is modules["growth_ranker.experience"].estimate_years_experience
    assert facade.score_candidate is modules["growth_ranker.scoring"].score_candidate
    assert facade.classify_score is modules["growth_ranker.scoring"].classify_score
    assert facade.run_analysis is modules["growth_ranker.exporting"].run_analysis


def _dashboard_rows():
    return [
        {"Nombre": "Laura", "Archivo": "laura.pdf", "Ajuste %": 94, "Clasificación": "GRUPO 1 - Prioridad alta"},
        {"Nombre": "Carlos", "Archivo": "carlos.docx", "Ajuste %": 85, "Clasificación": "GRUPO 2 - Entrevistar / validar"},
        {"Nombre": "Andrea", "Archivo": "andrea.pdf", "Ajuste %": 76, "Clasificación": "GRUPO 3 - Reserva con potencial"},
        {"Nombre": "Pedro", "Archivo": "pedro.pdf", "Ajuste %": 55, "Clasificación": "NO PRIORIZAR para Lead"},
        {"Nombre": "Scan", "Archivo": "scan.pdf", "Ajuste %": 0, "Clasificación": "REVISIÓN MANUAL - sin texto extraíble"},
        {"Nombre": "Broken", "Archivo": "broken.pdf", "Ajuste %": 0, "Clasificación": "ERROR"},
    ]


def test_dashboard_summary_separates_scored_candidates_from_manual_review():
    from growth_ranker.dashboard import dashboard_summary

    summary = dashboard_summary(_dashboard_rows())

    assert summary == {
        "total": 6,
        "scored": 4,
        "average_score": 77.5,
        "group_1": 1,
        "group_2": 1,
        "group_3": 1,
        "not_prioritized": 1,
        "review": 2,
    }


def test_score_histogram_uses_fixed_twenty_point_buckets_and_ignores_unscored_rows():
    from growth_ranker.dashboard import score_histogram

    histogram = score_histogram(_dashboard_rows())

    assert histogram == [
        ("0-19", 0),
        ("20-39", 0),
        ("40-59", 1),
        ("60-79", 1),
        ("80-100", 2),
    ]


def test_dashboard_filters_by_group_and_search_text_case_insensitively():
    from growth_ranker.dashboard import filter_rows

    rows = _dashboard_rows()
    assert [row["Nombre"] for row in filter_rows(rows, group="GRUPO 2")] == ["Carlos"]
    assert [row["Nombre"] for row in filter_rows(rows, query="ANDREA")] == ["Andrea"]
    assert [row["Nombre"] for row in filter_rows(rows, query="pdf", group="GRUPO 1")] == ["Laura"]


def test_top_candidates_returns_highest_scored_reviewable_rows_only():
    from growth_ranker.dashboard import top_candidates

    rows = _dashboard_rows()
    assert [row["Nombre"] for row in top_candidates(rows, limit=3)] == ["Laura", "Carlos", "Andrea"]


def test_phrase_hits_uses_token_boundaries_for_short_keywords():
    text = "Gestión de proyectos, productos y procesos para una empresa de tecnología."
    assert phrase_hits(text, ["pr"]) == []


def test_phrase_hits_matches_real_phrase_case_and_accents_independently():
    text = "Lideré la ESTRATEGIA DE MARKETING y la generación de demanda regional."
    hits = phrase_hits(text, ["estrategia de marketing", "generacion de demanda"])
    assert hits == ["estrategia de marketing", "generacion de demanda"]


def test_experience_estimation_ignores_education_ranges():
    text = """
    EDUCACIÓN
    Universidad Nacional - Administración de Empresas
    2014 - 2019

    EXPERIENCIA PROFESIONAL
    Growth Specialist - Empresa A
    Enero 2021 - Diciembre 2022
    Growth Manager - Empresa B
    Enero 2023 - Enero 2025
    """
    years = estimate_years_experience(text, now_year=2026, now_month=9)
    assert 3.8 <= years <= 4.2


def test_experience_estimation_merges_overlapping_jobs():
    text = """
    EXPERIENCIA LABORAL
    Marketing Manager | Enero 2020 - Diciembre 2023
    Consultor Growth | Enero 2022 - Diciembre 2024
    """
    years = estimate_years_experience(text, now_year=2026, now_month=9)
    assert 4.8 <= years <= 5.1


def test_budget_keyword_alone_is_not_treated_as_strong_management_evidence():
    weak = "Trabajé siguiendo el presupuesto aprobado por la dirección."
    strong = "Administré presupuesto anual de COP 800 millones para paid media y optimicé la inversión."
    weak_result = score_candidate(weak, years_override=7)
    strong_result = score_candidate(strong, years_override=7)
    assert strong_result["details"]["presupuesto"] > weak_result["details"]["presupuesto"]


def test_quantified_outcome_scores_more_than_bare_percentage():
    bare = "Disponibilidad 100% y dominio de herramientas digitales."
    outcome = "Incrementé las ventas 35%, reduje el CAC 22% y alcancé ROAS 4.1x."
    bare_result = score_candidate(bare, years_override=6)
    outcome_result = score_candidate(outcome, years_override=6)
    assert outcome_result["details"]["resultados_cuantitativos"] > bare_result["details"]["resultados_cuantitativos"]


def test_seniority_is_its_own_score_component_not_global_multiplier():
    text = """
    Growth marketing, estrategia de marketing, generación de demanda, ventas B2B,
    pipeline, ROAS, CAC, campañas 360, Meta Ads, Google Ads, HubSpot,
    liderazgo de equipo, e-commerce. Incrementé ventas 40% y administré COP 500 millones.
    """
    junior = score_candidate(text, years_override=3)
    senior = score_candidate(text, years_override=8)
    assert junior["details"]["experiencia"] < senior["details"]["experiencia"]
    for key in junior["details"]:
        if key != "experiencia":
            assert junior["details"][key] == senior["details"][key]


def test_classification_thresholds_are_stable():
    assert classify_score(90) == "GRUPO 1 - Prioridad alta"
    assert classify_score(82) == "GRUPO 2 - Entrevistar / validar"
    assert classify_score(74) == "GRUPO 3 - Reserva con potencial"
    assert classify_score(73.9) == "NO PRIORIZAR para Lead"


def test_end_to_end_folder_ranks_stronger_candidate_first_and_exports(tmp_path):
    from docx import Document
    from openpyxl import load_workbook
    from marketing_cv_ranker import run_analysis

    strong_doc = Document()
    strong_doc.add_paragraph("Laura Gómez")
    strong_doc.add_paragraph("EXPERIENCIA PROFESIONAL")
    strong_doc.add_paragraph("Growth Marketing Manager | Enero 2017 - Actualidad")
    strong_doc.add_paragraph(
        "Lideré estrategia de marketing y generación de demanda B2B; administré presupuesto "
        "de COP 800 millones en Meta Ads y Google Ads, usando HubSpot y coordinando equipo comercial."
    )
    strong_doc.add_paragraph("Incrementé ventas 42%, reduje CAC 25% y alcancé ROAS 4.5x en e-commerce.")
    strong_path = tmp_path / "Laura_Gomez.docx"
    strong_doc.save(strong_path)

    weak_doc = Document()
    weak_doc.add_paragraph("Pedro Pérez")
    weak_doc.add_paragraph("EDUCACIÓN")
    weak_doc.add_paragraph("Universidad 2014 - 2019")
    weak_doc.add_paragraph("EXPERIENCIA PROFESIONAL")
    weak_doc.add_paragraph("Community Manager | Enero 2024 - Actualidad")
    weak_doc.add_paragraph("Publicación de contenidos y apoyo en redes sociales.")
    weak_path = tmp_path / "Pedro_Perez.docx"
    weak_doc.save(weak_path)

    rows, xlsx_path, csv_path = run_analysis(tmp_path)

    assert len(rows) == 2
    assert rows[0]["Nombre"] == "Laura Gómez"
    assert rows[0]["Ajuste %"] > rows[1]["Ajuste %"]
    assert xlsx_path.exists()
    assert csv_path.exists()

    workbook = load_workbook(xlsx_path)
    sheet = workbook["Ranking Marketing"]
    assert sheet["A2"].value == 1
    assert sheet["B2"].value == "Laura Gómez"
