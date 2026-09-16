#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Ranker local de hojas de vida para perfiles Growth & Marketing Lead.

- Procesa PDF y DOCX de una carpeta.
- Extrae nombre, teléfono y correo.
- Estima experiencia laboral evitando, cuando es posible, periodos de educación.
- Puntúa evidencia de marketing con coincidencia por palabras/frases completas.
- Premia resultados cuantificados contextualizados.
- Exporta ranking a Excel y CSV.
- Incluye GUI Tkinter para uso como aplicación de escritorio.
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import subprocess
import sys
import threading
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable

from pypdf import PdfReader
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


APP_TITLE = "Ranking de Candidatos - Marketing"
PROFILE = {
    "cargo": "Growth & Marketing Lead",
    "experiencia_objetivo_min": 6,
    "experiencia_objetivo_ideal": 8,
    "criterios": {
        "estrategia_growth": {
            "peso": 15,
            "keywords": [
                "growth marketing", "growth", "estrategia de marketing",
                "marketing estrategico", "planeacion estrategica", "go to market",
                "go-to-market", "posicionamiento", "generacion de demanda",
                "demand generation", "expansion de mercado", "market expansion",
                "desarrollo de negocio", "business development",
            ],
        },
        "conexion_comercial": {
            "peso": 12,
            "keywords": [
                "ventas", "comercial", "equipo comercial", "sales", "pipeline",
                "funnel", "embudo", "lead", "leads", "conversion", "cierre",
                "clientes b2b", "b2b", "b2c", "revenue", "facturacion",
            ],
        },
        "presupuesto": {
            "peso": 9,
            "keywords": [
                "presupuesto", "budget", "ad spend", "inversion",
                "manejo de presupuesto", "gestion presupuestal", "paid media budget",
            ],
            "strong_patterns": [
                r"\b(?:administre|administro|gestione|gestiono|maneje|manejo|lidere|lidero|responsable de)\b.{0,45}\b(?:presupuesto|budget|inversion|ad spend)\b",
                r"\b(?:presupuesto|budget|inversion|ad spend)\b.{0,45}\b(?:de|por|superior a|hasta)?\s*(?:cop|usd|\$)\s*[\d\.,]+",
                r"\b(?:cop|usd|\$)\s*[\d\.,]+\s*(?:millones|mm|m)?\b.{0,45}\b(?:presupuesto|budget|inversion|ad spend|paid media)\b",
            ],
        },
        "metricas_resultados": {
            "peso": 10,
            "keywords": [
                "roi", "roas", "cac", "ltv", "cpl", "cpa", "ctr", "kpi", "kpis",
                "tasa de conversion", "conversion rate", "ventas generadas",
                "crecimiento de ingresos", "performance", "rentabilidad", "retorno",
            ],
        },
        "campanas_activaciones_eventos": {
            "peso": 7,
            "keywords": [
                "campanas 360", "campana 360", "atl", "btl", "activaciones",
                "activacion de marca", "eventos", "lanzamientos", "ferias",
                "networking", "relaciones publicas", "public relations", "pr",
            ],
        },
        "digital_paid_media": {
            "peso": 7,
            "keywords": [
                "meta ads", "facebook ads", "google ads", "tiktok ads",
                "linkedin ads", "paid media", "seo", "sem", "google analytics",
                "ga4", "performance marketing",
            ],
        },
        "crm_automatizacion": {
            "peso": 4,
            "keywords": [
                "hubspot", "salesforce", "crm", "automatizacion", "automation",
                "mailchimp", "manychat", "zoho", "kommo", "bitrix24",
                "marketing automation",
            ],
        },
        "liderazgo": {
            "peso": 4,
            "keywords": [
                "lidere equipo", "liderazgo", "coordine equipo", "coordinacion de equipos",
                "marketing manager", "growth manager", "director de marketing",
                "jefe de mercadeo", "lider de marketing", "agencias", "proveedores",
                "stakeholders",
            ],
        },
        "sectores_deseables": {
            "peso": 2,
            "keywords": [
                "e-commerce", "ecommerce", "retail", "logistica", "logistics",
                "startup", "marketplace", "fulfillment", "comercio electronico",
            ],
        },
    },
}

EMAIL_RE = re.compile(r"\b[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}\b")
PHONE_RE = re.compile(r"(?<!\d)(?:\+?57[\s\-\.]?)?(?:3\d{2})[\s\-\.]?\d{3}[\s\-\.]?\d{4}(?!\d)")

MONTHS = {
    "enero": 1, "ene": 1, "january": 1, "jan": 1,
    "febrero": 2, "feb": 2, "february": 2,
    "marzo": 3, "mar": 3, "march": 3,
    "abril": 4, "abr": 4, "april": 4, "apr": 4,
    "mayo": 5, "may": 5,
    "junio": 6, "jun": 6, "june": 6,
    "julio": 7, "jul": 7, "july": 7,
    "agosto": 8, "ago": 8, "august": 8, "aug": 8,
    "septiembre": 9, "sep": 9, "sept": 9, "september": 9,
    "octubre": 10, "oct": 10, "october": 10,
    "noviembre": 11, "nov": 11, "november": 11,
    "diciembre": 12, "dic": 12, "december": 12, "dec": 12,
}

WORK_HEADINGS = {
    "experiencia", "experiencia laboral", "experiencia profesional",
    "trayectoria profesional", "historial laboral", "work experience",
    "professional experience", "employment history", "career experience",
}
EDUCATION_HEADINGS = {
    "educacion", "formacion", "formacion academica", "educacion academica",
    "estudios", "academic background", "education", "certificaciones",
    "certifications", "cursos", "courses",
}
OTHER_SECTION_HEADINGS = {
    "perfil", "perfil profesional", "resumen", "resumen profesional", "habilidades",
    "skills", "competencias", "idiomas", "languages", "referencias", "references",
    "proyectos", "projects", "contacto", "contact",
}


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFKD", text or "")
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.lower().replace("–", "-").replace("—", "-")
    return re.sub(r"\s+", " ", text).strip()


def _normalize_line(text: str) -> str:
    return re.sub(r"[^a-z0-9 +&/\-]", "", normalize_text(text)).strip(" :-|")


def phrase_hits(text: str, phrases: Iterable[str]) -> list[str]:
    norm = normalize_text(text)
    hits: list[str] = []
    for phrase in phrases:
        phrase_norm = normalize_text(phrase)
        if not phrase_norm:
            continue
        pattern = rf"(?<!\w){re.escape(phrase_norm)}(?!\w)"
        if re.search(pattern, norm, re.IGNORECASE):
            hits.append(phrase)
    return hits


def read_pdf(path: Path) -> str:
    chunks: list[str] = []
    reader = PdfReader(str(path))
    for page in reader.pages:
        try:
            chunks.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(chunks)


def read_docx(path: Path) -> str:
    doc = Document(str(path))
    chunks: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            chunks.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(chunks)


def extract_text(path: Path) -> str:
    if path.suffix.lower() == ".pdf":
        return read_pdf(path)
    if path.suffix.lower() == ".docx":
        return read_docx(path)
    return ""


def normalize_phone(phone: str) -> str:
    digits = re.sub(r"\D", "", phone)
    if digits.startswith("57") and len(digits) >= 12:
        digits = digits[2:]
    if len(digits) == 10:
        return f"{digits[:3]} {digits[3:6]} {digits[6:]}"
    return phone.strip()


def extract_email(text: str) -> str:
    matches = EMAIL_RE.findall(text)
    return matches[0] if matches else ""


def extract_phone(text: str) -> str:
    phones: list[str] = []
    for match in PHONE_RE.findall(text):
        phone = normalize_phone(match)
        if phone not in phones:
            phones.append(phone)
    return " / ".join(phones[:2])


def guess_name(filename: str, text: str) -> str:
    lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines() if line.strip()]
    bad_tokens = {
        "perfil", "experiencia", "educacion", "contacto", "curriculum", "hoja de vida",
        "resumen profesional", "habilidades", "formacion", "marketing", "growth",
    }
    for line in lines[:12]:
        low = normalize_text(line)
        if any(token in low for token in bad_tokens):
            continue
        words = line.split()
        alpha_ratio = sum(ch.isalpha() or ch.isspace() for ch in line) / max(1, len(line))
        if 2 <= len(words) <= 6 and alpha_ratio > 0.80 and len(line) <= 60:
            return line.title()
    stem = Path(filename).stem
    stem = re.sub(r"(?i)^cv[\s_\-]*", "", stem)
    stem = re.sub(r"[_\-]+", " ", stem)
    return re.sub(r"\s+", " ", stem).strip().title()


def _looks_like_heading(line: str) -> bool:
    stripped = line.strip()
    if not stripped or len(stripped) > 70:
        return False
    words = stripped.split()
    if len(words) > 7:
        return False
    letters = [ch for ch in stripped if ch.isalpha()]
    if not letters:
        return False
    uppercase_ratio = sum(ch.isupper() for ch in letters) / len(letters)
    return uppercase_ratio >= 0.72 or stripped.endswith(":")


def extract_work_experience_text(text: str) -> str:
    lines = text.splitlines()
    mode = "neutral"
    work_found = False
    work_lines: list[str] = []
    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue
        heading = _normalize_line(line)
        if heading in WORK_HEADINGS or any(heading.startswith(h + " ") for h in WORK_HEADINGS if len(h) > 10):
            mode = "work"
            work_found = True
            continue
        if heading in EDUCATION_HEADINGS or any(heading.startswith(h + " ") for h in EDUCATION_HEADINGS if len(h) > 10):
            mode = "education"
            continue
        if heading in OTHER_SECTION_HEADINGS and mode == "work":
            mode = "neutral"
            continue
        if mode == "work":
            if _looks_like_heading(line) and not re.search(r"\b(?:19|20)\d{2}\b", line):
                norm = _normalize_line(line)
                if norm not in WORK_HEADINGS:
                    mode = "neutral"
                    continue
            work_lines.append(line)

    if work_found and work_lines:
        return "\n".join(work_lines)

    filtered: list[str] = []
    education_mode = False
    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue
        heading = _normalize_line(line)
        if heading in EDUCATION_HEADINGS:
            education_mode = True
            continue
        if heading in WORK_HEADINGS or (education_mode and heading in OTHER_SECTION_HEADINGS):
            education_mode = False
        if education_mode:
            continue
        if re.search(r"\b(?:universidad|university|instituto|college|licenciatura|pregrado|posgrado|maestria|master|diplomado)\b", normalize_text(line)):
            continue
        filtered.append(line)
    return "\n".join(filtered)


def month_num(token: str) -> int:
    return MONTHS.get(normalize_text(token).strip("."), 1)


def extract_date_ranges(text: str, now: datetime | None = None) -> list[tuple[datetime, datetime]]:
    now = now or datetime.now()
    norm = normalize_text(text)
    ranges: list[tuple[datetime, datetime]] = []
    month_names = "|".join(sorted((re.escape(k) for k in MONTHS), key=len, reverse=True))
    present_words = r"(?:actualidad|actualmente|presente|present|current|hoy)"
    month_pattern = re.compile(
        rf"(?<!\w)({month_names})\s+((?:19|20)\d{{2}})\s*-\s*"
        rf"(?:(?:({month_names})\s+((?:19|20)\d{{2}}))|({present_words}))(?!\w)",
        re.IGNORECASE,
    )
    occupied: list[tuple[int, int]] = []
    for match in month_pattern.finditer(norm):
        sm, sy, em, ey, present = match.groups()
        start = datetime(int(sy), month_num(sm), 1)
        end = now if present else datetime(int(ey), month_num(em), 1)
        if 1990 <= start.year <= now.year and start <= end:
            ranges.append((start, end))
            occupied.append(match.span())
    year_pattern = re.compile(
        rf"(?<!\w)((?:19|20)\d{{2}})\s*-\s*((?:19|20)\d{{2}}|actualidad|actualmente|presente|present|current)(?!\w)",
        re.IGNORECASE,
    )
    for match in year_pattern.finditer(norm):
        if any(match.start() >= s and match.end() <= e for s, e in occupied):
            continue
        sy, ey = match.groups()
        start = datetime(int(sy), 1, 1)
        end = now if not ey.isdigit() else datetime(int(ey), 12, 1)
        if 1990 <= start.year <= now.year and start <= end:
            ranges.append((start, end))
    return ranges


def merge_intervals(intervals: list[tuple[datetime, datetime]]) -> list[tuple[datetime, datetime]]:
    if not intervals:
        return []
    ordered = sorted(intervals, key=lambda item: item[0])
    merged: list[list[datetime]] = [[ordered[0][0], ordered[0][1]]]
    for start, end in ordered[1:]:
        last = merged[-1]
        if (start - last[1]).days <= 31:
            last[1] = max(last[1], end)
        else:
            merged.append([start, end])
    return [(start, end) for start, end in merged]


def estimate_years_experience(text: str, *, now_year: int | None = None, now_month: int | None = None) -> float:
    now = datetime(now_year or datetime.now().year, now_month or datetime.now().month, 1)
    work_text = extract_work_experience_text(text)
    intervals = merge_intervals(extract_date_ranges(work_text, now=now))
    total_days = sum((end - start).days for start, end in intervals)
    return round(total_days / 365.25, 1)


def find_money_mentions(text: str) -> list[str]:
    patterns = [
        r"\$\s?[\d\.,]+\s?(?:mm|m|millones|cop|usd)?",
        r"\b\d+(?:[\.,]\d+)?\s?(?:millones|mm)\s?(?:cop|usd)?\b",
        r"\b(?:cop|usd)\s?[\d\.,]+(?:\s?(?:millones|mm|m))?",
    ]
    found: list[str] = []
    for pattern in patterns:
        found.extend(re.findall(pattern, text, re.IGNORECASE))
    return list(dict.fromkeys(found))[:5]


def find_percentage_mentions(text: str) -> list[str]:
    return list(dict.fromkeys(re.findall(r"\b\d+(?:[\.,]\d+)?\s?%", text)))[:8]


def find_ratio_mentions(text: str) -> list[str]:
    return list(dict.fromkeys(re.findall(r"\b\d+(?:[\.,]\d+)?\s?x\b", text, re.IGNORECASE)))[:5]


def _contextual_quantified_results(text: str) -> list[str]:
    norm = normalize_text(text)
    metric = r"(?:ventas|ingresos|revenue|conversion|roi|roas|cac|ltv|cpl|cpa|ctr|leads?|facturacion|trafico|retencion)"
    verb = r"(?:aumente|aumento|incremente|incremento|creci|crecimiento|reduje|reduccion|disminui|mejore|mejora|optimice|optimizo|alcance|logre|genere|duplique|triplique)"
    value = r"(?:\d+(?:[\.,]\d+)?\s?%|\d+(?:[\.,]\d+)?\s?x|(?:cop|usd|\$)\s?[\d\.,]+(?:\s?(?:millones|mm|m))?)"
    patterns = [
        rf"\b{verb}\b.{{0,55}}\b{metric}\b.{{0,35}}{value}",
        rf"\b{verb}\b.{{0,55}}{value}.{{0,35}}\b{metric}\b",
        rf"\b{metric}\b.{{0,35}}{value}",
    ]
    found: list[str] = []
    for pattern in patterns:
        for match in re.finditer(pattern, norm, re.IGNORECASE):
            snippet = match.group(0).strip()
            if snippet not in found:
                found.append(snippet)
    return found[:8]


def _coverage_ratio(hit_count: int) -> float:
    if hit_count <= 0:
        return 0.0
    if hit_count == 1:
        return 0.40
    if hit_count == 2:
        return 0.70
    if hit_count == 3:
        return 0.90
    return 1.0


def experience_points(years: float) -> float:
    if years >= 8:
        return 15.0
    if years >= 6:
        return 13.5
    if years >= 4:
        return 10.0
    if years >= 2:
        return 6.0
    if years > 0:
        return 3.0
    return 0.0


def classify_score(score: float) -> str:
    if score >= 90:
        return "GRUPO 1 - Prioridad alta"
    if score >= 82:
        return "GRUPO 2 - Entrevistar / validar"
    if score >= 74:
        return "GRUPO 3 - Reserva con potencial"
    return "NO PRIORIZAR para Lead"


def score_candidate(text: str, *, years_override: float | None = None) -> dict:
    norm = normalize_text(text)
    years = years_override if years_override is not None else estimate_years_experience(text)
    evidence: dict[str, list[str]] = {}
    details: dict[str, float] = {}
    total = 0.0
    for criterion, cfg in PROFILE["criterios"].items():
        hits = phrase_hits(norm, cfg["keywords"])
        evidence[criterion] = hits
        ratio = _coverage_ratio(len(hits))
        strong_patterns = cfg.get("strong_patterns", [])
        strong_hits = [pattern for pattern in strong_patterns if re.search(pattern, norm, re.IGNORECASE)]
        if strong_hits:
            ratio = max(ratio, 0.85 if len(strong_hits) == 1 else 1.0)
        points = round(cfg["peso"] * ratio, 1)
        details[criterion] = points
        total += points
    exp_points = experience_points(years)
    details["experiencia"] = exp_points
    total += exp_points
    money = find_money_mentions(text)
    percentages = find_percentage_mentions(text)
    ratios = find_ratio_mentions(text)
    quantified = _contextual_quantified_results(text)
    quantitative_points = 0.0
    if money or percentages or ratios:
        quantitative_points = 1.5
    if quantified:
        quantitative_points = min(15.0, 3.0 + len(quantified) * 4.0)
    details["resultados_cuantitativos"] = round(quantitative_points, 1)
    total += quantitative_points
    raw = min(100.0, total)
    if not evidence["estrategia_growth"] or not evidence["conexion_comercial"]:
        raw = min(raw, 74.0)
    if not evidence["presupuesto"] and not evidence["metricas_resultados"]:
        raw = min(raw, 79.0)
    score = round(raw, 1)
    return {
        "score": score,
        "recommendation": classify_score(score),
        "years": round(float(years), 1),
        "evidence": evidence,
        "details": details,
        "money": money,
        "percentages": percentages,
        "ratios": ratios,
        "quantified_results": quantified,
    }


def build_experience_summary(result: dict) -> str:
    parts: list[str] = []
    if result["years"]:
        parts.append(f"Experiencia estimada: {result['years']} años")
    labels = {
        "estrategia_growth": "estrategia/growth",
        "conexion_comercial": "marketing + comercial",
        "presupuesto": "presupuesto",
        "metricas_resultados": "métricas/ROI",
        "campanas_activaciones_eventos": "campañas/eventos",
        "digital_paid_media": "paid media",
        "crm_automatizacion": "CRM/automatización",
        "liderazgo": "liderazgo",
        "sectores_deseables": "sectores afines",
    }
    strong = [labels[key] for key, hits in result["evidence"].items() if len(hits) >= 2]
    if strong:
        parts.append("Fortalezas: " + ", ".join(strong))
    if result["quantified_results"]:
        parts.append(f"Resultados cuantificados detectados: {len(result['quantified_results'])}")
    if result["money"]:
        parts.append("Montos: " + ", ".join(result["money"][:3]))
    return " | ".join(parts)


def build_gaps(result: dict) -> str:
    missing: list[str] = []
    if not result["evidence"]["estrategia_growth"]:
        missing.append("sin evidencia clara de estrategia/growth")
    if not result["evidence"]["conexion_comercial"]:
        missing.append("sin conexión comercial clara")
    if not result["evidence"]["presupuesto"]:
        missing.append("sin evidencia clara de presupuesto")
    if not result["evidence"]["metricas_resultados"]:
        missing.append("sin evidencia fuerte de ROI/métricas")
    if not result["evidence"]["campanas_activaciones_eventos"]:
        missing.append("poca evidencia de eventos/activaciones")
    if result["years"] < PROFILE["experiencia_objetivo_min"]:
        missing.append("seniority menor a 6 años")
    if not result["quantified_results"]:
        missing.append("sin resultados cuantificados contextualizados")
    return "; ".join(missing)


def build_score_breakdown(result: dict) -> str:
    labels = {
        "estrategia_growth": "Estrategia",
        "conexion_comercial": "Comercial",
        "presupuesto": "Presupuesto",
        "metricas_resultados": "Métricas",
        "campanas_activaciones_eventos": "Campañas",
        "digital_paid_media": "Digital",
        "crm_automatizacion": "CRM",
        "liderazgo": "Liderazgo",
        "sectores_deseables": "Sector",
        "experiencia": "Experiencia",
        "resultados_cuantitativos": "Resultados",
    }
    return " | ".join(f"{labels[k]} {v:g}" for k, v in result["details"].items())


def process_folder(folder: Path, progress: Callable[[int, int, str], None] | None = None) -> list[dict]:
    files = sorted(path for path in folder.iterdir() if path.is_file() and path.suffix.lower() in {".pdf", ".docx"})
    rows: list[dict] = []
    total = len(files)
    for index, path in enumerate(files, 1):
        if progress:
            progress(index - 1, total, path.name)
        try:
            text = extract_text(path)
            if not text.strip():
                rows.append({
                    "Archivo": path.name,
                    "Nombre": guess_name(path.name, ""),
                    "Celular": "",
                    "Correo": "",
                    "Años experiencia estimados": 0,
                    "Ajuste %": 0,
                    "Clasificación": "REVISIÓN MANUAL - sin texto extraíble",
                    "Experiencia / fortalezas": "",
                    "Brechas": "El archivo no contiene texto extraíble.",
                    "Desglose": "",
                })
                continue
            result = score_candidate(text)
            rows.append({
                "Archivo": path.name,
                "Nombre": guess_name(path.name, text),
                "Celular": extract_phone(text),
                "Correo": extract_email(text),
                "Años experiencia estimados": result["years"],
                "Ajuste %": result["score"],
                "Clasificación": result["recommendation"],
                "Experiencia / fortalezas": build_experience_summary(result),
                "Brechas": build_gaps(result),
                "Desglose": build_score_breakdown(result),
            })
        except Exception as exc:
            rows.append({
                "Archivo": path.name,
                "Nombre": path.stem,
                "Celular": "",
                "Correo": "",
                "Años experiencia estimados": 0,
                "Ajuste %": 0,
                "Clasificación": "ERROR",
                "Experiencia / fortalezas": "",
                "Brechas": str(exc),
                "Desglose": "",
            })
    rows.sort(key=lambda row: row["Ajuste %"], reverse=True)
    if progress:
        progress(total, total, "Completado")
    return rows


def export_csv(rows: list[dict], out_path: Path) -> None:
    if not rows:
        return
    with out_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def export_excel(rows: list[dict], out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ranking Marketing"
    headers = [
        "Ranking", "Nombre", "Celular", "Correo", "Años experiencia estimados",
        "Ajuste %", "Clasificación", "Experiencia / fortalezas", "Brechas",
        "Desglose", "Archivo",
    ]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="145E77")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for ranking, row in enumerate(rows, 1):
        ws.append([
            ranking, row["Nombre"], row["Celular"], row["Correo"],
            row["Años experiencia estimados"], row["Ajuste %"], row["Clasificación"],
            row["Experiencia / fortalezas"], row["Brechas"], row["Desglose"], row["Archivo"],
        ])
    for r in range(2, ws.max_row + 1):
        cls = str(ws.cell(r, 7).value or "")
        if "GRUPO 1" in cls:
            color = "D9EAD3"
        elif "GRUPO 2" in cls:
            color = "FFF2CC"
        elif "GRUPO 3" in cls:
            color = "FCE5CD"
        else:
            color = "F4CCCC"
        for c in range(1, ws.max_column + 1):
            ws.cell(r, c).fill = PatternFill("solid", fgColor=color)
            ws.cell(r, c).alignment = Alignment(vertical="top", wrap_text=True)
    widths = {1: 9, 2: 28, 3: 18, 4: 34, 5: 22, 6: 12, 7: 28, 8: 62, 9: 48, 10: 70, 11: 35}
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(out_path)


def run_analysis(folder: Path, output_dir: Path | None = None, progress=None) -> tuple[list[dict], Path, Path]:
    output_dir = output_dir or folder
    output_dir.mkdir(parents=True, exist_ok=True)
    rows = process_folder(folder, progress=progress)
    xlsx_path = output_dir / "ranking_candidatos_marketing.xlsx"
    csv_path = output_dir / "ranking_candidatos_marketing.csv"
    export_excel(rows, xlsx_path)
    export_csv(rows, csv_path)
    return rows, xlsx_path, csv_path


def open_path(path: Path) -> None:
    if sys.platform.startswith("win"):
        os.startfile(str(path))  # type: ignore[attr-defined]
    elif sys.platform == "darwin":
        subprocess.Popen(["open", str(path)])
    else:
        subprocess.Popen(["xdg-open", str(path)])


def launch_gui() -> None:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
    root = tk.Tk()
    root.title(APP_TITLE)
    root.geometry("1080x680")
    root.minsize(920, 560)
    selected_folder = tk.StringVar()
    status = tk.StringVar(value="Selecciona una carpeta con hojas de vida PDF o DOCX.")
    progress_value = tk.DoubleVar(value=0)
    last_excel: dict[str, Path | None] = {"path": None}
    outer = ttk.Frame(root, padding=18)
    outer.pack(fill="both", expand=True)
    ttk.Label(outer, text="Ranking de candidatos - Marketing", font=("Segoe UI", 18, "bold")).pack(anchor="w")
    ttk.Label(outer, text="Growth & Marketing Lead · análisis local · PDF y DOCX", font=("Segoe UI", 10)).pack(anchor="w", pady=(2, 16))
    picker = ttk.Frame(outer)
    picker.pack(fill="x")
    entry = ttk.Entry(picker, textvariable=selected_folder)
    entry.pack(side="left", fill="x", expand=True)

    def choose_folder() -> None:
        folder = filedialog.askdirectory(title="Selecciona la carpeta de hojas de vida")
        if folder:
            selected_folder.set(folder)

    ttk.Button(picker, text="Seleccionar carpeta", command=choose_folder).pack(side="left", padx=(8, 0))
    button_row = ttk.Frame(outer)
    button_row.pack(fill="x", pady=(12, 8))
    analyze_btn = ttk.Button(button_row, text="Analizar candidatos")
    analyze_btn.pack(side="left")
    open_excel_btn = ttk.Button(button_row, text="Abrir Excel", state="disabled")
    open_excel_btn.pack(side="left", padx=(8, 0))
    open_folder_btn = ttk.Button(button_row, text="Abrir carpeta", state="disabled")
    open_folder_btn.pack(side="left", padx=(8, 0))
    progress_bar = ttk.Progressbar(outer, maximum=100, variable=progress_value)
    progress_bar.pack(fill="x", pady=(4, 4))
    ttk.Label(outer, textvariable=status).pack(anchor="w", pady=(0, 10))
    columns = ("rank", "name", "score", "years", "group", "file")
    tree = ttk.Treeview(outer, columns=columns, show="headings", height=18)
    headings = {"rank": "#", "name": "Candidato", "score": "Ajuste %", "years": "Años exp.", "group": "Clasificación", "file": "Archivo"}
    widths = {"rank": 45, "name": 220, "score": 85, "years": 85, "group": 250, "file": 260}
    for col in columns:
        tree.heading(col, text=headings[col])
        tree.column(col, width=widths[col], anchor="center" if col in {"rank", "score", "years"} else "w")
    tree.pack(fill="both", expand=True)

    def refresh_table(rows: list[dict]) -> None:
        for item in tree.get_children():
            tree.delete(item)
        for rank, row in enumerate(rows, 1):
            tree.insert("", "end", values=(rank, row["Nombre"], row["Ajuste %"], row["Años experiencia estimados"], row["Clasificación"], row["Archivo"]))

    def start_analysis() -> None:
        folder_text = selected_folder.get().strip()
        folder = Path(folder_text) if folder_text else None
        if not folder or not folder.exists() or not folder.is_dir():
            messagebox.showwarning(APP_TITLE, "Selecciona una carpeta válida.")
            return
        candidates = [p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in {".pdf", ".docx"}]
        if not candidates:
            messagebox.showwarning(APP_TITLE, "La carpeta no contiene archivos PDF o DOCX.")
            return
        analyze_btn.configure(state="disabled")
        open_excel_btn.configure(state="disabled")
        open_folder_btn.configure(state="disabled")
        progress_value.set(0)
        status.set(f"Preparando {len(candidates)} hojas de vida...")

        def update_progress(done: int, total: int, name: str) -> None:
            pct = (done / total * 100) if total else 0
            root.after(0, lambda: progress_value.set(pct))
            root.after(0, lambda: status.set(f"Procesando {done}/{total}: {name}"))

        def worker() -> None:
            try:
                rows, xlsx, _csv = run_analysis(folder, progress=update_progress)
                last_excel["path"] = xlsx
                root.after(0, lambda: refresh_table(rows))
                root.after(0, lambda: status.set(f"Listo: {len(rows)} hojas de vida procesadas. Excel guardado en la carpeta seleccionada."))
                root.after(0, lambda: progress_value.set(100))
                root.after(0, lambda: open_excel_btn.configure(state="normal"))
                root.after(0, lambda: open_folder_btn.configure(state="normal"))
            except Exception as exc:
                root.after(0, lambda: messagebox.showerror(APP_TITLE, f"No se pudo completar el análisis:\n{exc}"))
                root.after(0, lambda: status.set("El análisis terminó con error."))
            finally:
                root.after(0, lambda: analyze_btn.configure(state="normal"))

        threading.Thread(target=worker, daemon=True).start()

    analyze_btn.configure(command=start_analysis)
    open_excel_btn.configure(command=lambda: open_path(last_excel["path"]) if last_excel["path"] else None)
    open_folder_btn.configure(command=lambda: open_path(Path(selected_folder.get())) if selected_folder.get() else None)
    root.mainloop()


def cli_main(args: argparse.Namespace) -> int:
    folder = Path(args.carpeta).expanduser().resolve()
    if not folder.exists() or not folder.is_dir():
        print(f"No existe la carpeta: {folder}", file=sys.stderr)
        return 2
    out = Path(args.salida).expanduser().resolve() if args.salida else folder
    rows, xlsx, csv_path = run_analysis(folder, out)
    print("\n=== FILTRO TERMINADO ===")
    print(f"CVs procesadas: {len(rows)}")
    print(f"Excel: {xlsx}")
    print(f"CSV:   {csv_path}")
    print("\nTOP 10:")
    for row in rows[:10]:
        print(f"{row['Ajuste %']:>5}% | {row['Nombre'][:35]:35} | {row['Clasificación']}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="Ranking local de CVs para Growth & Marketing Lead")
    parser.add_argument("--carpeta", help="Carpeta con archivos PDF/DOCX. Si se omite, abre la interfaz gráfica.")
    parser.add_argument("--salida", help="Carpeta de salida; por defecto usa la carpeta de CVs.")
    args = parser.parse_args()
    if args.carpeta:
        return cli_main(args)
    launch_gui()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
