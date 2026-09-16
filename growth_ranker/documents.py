"""Document readers for supported CV formats."""

from pathlib import Path

from docx import Document
from pypdf import PdfReader


def read_pdf(path: Path) -> str:
    chunks: list[str] = []
    reader = PdfReader(str(path))
    for page in reader.pages:
        try:
            chunks.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(chunks)


def read_docx(path: Path) -> str:
    doc = Document(str(path))
    chunks: list[str] = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            chunks.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(chunks)


def extract_text(path: Path) -> str:
    if path.suffix.lower() == ".pdf":
        return read_pdf(path)
    if path.suffix.lower() == ".docx":
        return read_docx(path)
    return ""
