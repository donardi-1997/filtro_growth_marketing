"""Folder processing and Excel/CSV export."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .documents import extract_text
from .scoring import build_experience_summary, build_gaps, build_score_breakdown, score_candidate
from .text import extract_email, extract_phone, guess_name


def process_folder(folder: Path, progress: Callable[[int, int, str], None] | None = None) -> list[dict]:
    files = sorted(path for path in folder.iterdir() if path.is_file() and path.suffix.lower() in {".pdf", ".docx"})
    rows: list[dict] = []
    total = len(files)
    for index, path in enumerate(files, 1):
        if progress:
            progress(index - 1, total, path.name)
        try:
            text = extract_text(path)
            if not text.strip():
                rows.append({
                    "Archivo": path.name,
                    "Nombre": guess_name(path.name, ""),
                    "Celular": "",
                    "Correo": "",
                    "Años experiencia estimados": 0,
                    "Ajuste %": 0,
                    "Clasificación": "REVISIÓN MANUAL - sin texto extraíble",
                    "Experiencia / fortalezas": "",
                    "Brechas": "El archivo no contiene texto extraíble.",
                    "Desglose": "",
                })
                continue
            result = score_candidate(text)
            rows.append({
                "Archivo": path.name,
                "Nombre": guess_name(path.name, text),
                "Celular": extract_phone(text),
                "Correo": extract_email(text),
                "Años experiencia estimados": result["years"],
                "Ajuste %": result["score"],
                "Clasificación": result["recommendation"],
                "Experiencia / fortalezas": build_experience_summary(result),
                "Brechas": build_gaps(result),
                "Desglose": build_score_breakdown(result),
            })
        except Exception as exc:
            rows.append({
                "Archivo": path.name,
                "Nombre": path.stem,
                "Celular": "",
                "Correo": "",
                "Años experiencia estimados": 0,
                "Ajuste %": 0,
                "Clasificación": "ERROR",
                "Experiencia / fortalezas": "",
                "Brechas": str(exc),
                "Desglose": "",
            })
    rows.sort(key=lambda row: row["Ajuste %"], reverse=True)
    if progress:
        progress(total, total, "Completado")
    return rows


def export_csv(rows: list[dict], out_path: Path) -> None:
    if not rows:
        return
    with out_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def export_excel(rows: list[dict], out_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Ranking Marketing"
    headers = [
        "Ranking", "Nombre", "Celular", "Correo", "Años experiencia estimados",
        "Ajuste %", "Clasificación", "Experiencia / fortalezas", "Brechas",
        "Desglose", "Archivo",
    ]
    worksheet.append(headers)
    header_fill = PatternFill("solid", fgColor="145E77")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for ranking, row in enumerate(rows, 1):
        worksheet.append([
            ranking, row["Nombre"], row["Celular"], row["Correo"],
            row["Años experiencia estimados"], row["Ajuste %"], row["Clasificación"],
            row["Experiencia / fortalezas"], row["Brechas"], row["Desglose"], row["Archivo"],
        ])

    for row_number in range(2, worksheet.max_row + 1):
        classification = str(worksheet.cell(row_number, 7).value or "")
        if "GRUPO 1" in classification:
            color = "D9EAD3"
        elif "GRUPO 2" in classification:
            color = "FFF2CC"
        elif "GRUPO 3" in classification:
            color = "FCE5CD"
        else:
            color = "F4CCCC"
        for column in range(1, worksheet.max_column + 1):
            worksheet.cell(row_number, column).fill = PatternFill("solid", fgColor=color)
            worksheet.cell(row_number, column).alignment = Alignment(vertical="top", wrap_text=True)

    widths = {1: 9, 2: 28, 3: 18, 4: 34, 5: 22, 6: 12, 7: 28, 8: 62, 9: 48, 10: 70, 11: 35}
    for column, width in widths.items():
        worksheet.column_dimensions[get_column_letter(column)].width = width
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    workbook.save(out_path)


def run_analysis(folder: Path, output_dir: Path | None = None, progress=None) -> tuple[list[dict], Path, Path]:
    output_dir = output_dir or folder
    output_dir.mkdir(parents=True, exist_ok=True)
    rows = process_folder(folder, progress=progress)
    xlsx_path = output_dir / "ranking_candidatos_marketing.xlsx"
    csv_path = output_dir / "ranking_candidatos_marketing.csv"
    export_excel(rows, xlsx_path)
    export_csv(rows, csv_path)
    return rows, xlsx_path, csv_path
